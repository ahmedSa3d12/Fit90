import datetime
import os
import re
import sys
import unicodedata
from openpyxl import load_workbook

ROOT = r'E:\final_projects\asmaa\23-8-2026\FIT90_19-8'
SOURCE = os.path.join(ROOT, 'outputs', 'data-migration', 'final-upload-package-20260908', 'excluded_members_for_client.xlsx')
DEFAULT_OUTPUT = os.path.join(ROOT, 'outputs', 'data-migration', 'final-upload-package-20260908', 'add_excluded_members_except_duplicate_phones.sql')
EXPECTED_DB = sys.argv[1] if len(sys.argv) > 1 else 'metacodecx_fit90'
OUTPUT = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT
DRY_RUN = len(sys.argv) > 3 and sys.argv[3] == '--dry-run'

TRANS = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789')


def text(value):
    if value is None:
        return ''
    return unicodedata.normalize('NFKC', str(value)).replace('\u200b', '').replace('\ufeff', '').strip()


def sql_text(value):
    value = text(value)
    if not value:
        return 'NULL'
    value = value.replace('\\', '\\\\').replace("'", "''").replace('\x00', '').replace('\r', '\\r').replace('\n', '\\n')
    return f"'{value}'"


def date_text(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime('%Y-%m-%d')
    return text(value)[:10]


def insert_rows(rows, batch_size=250):
    statements = []
    for start in range(0, len(rows), batch_size):
        values = ',\n'.join('(' + ','.join(row) + ')' for row in rows[start:start + batch_size])
        statements.append(
            'INSERT INTO `_fit90_additional_member_import` '
            '(`legacy_member_id`, `name`, `phone`, `email`, `gender`, `card_number`, `date_of_birth`, `notes`) VALUES\n'
            + values + ';'
        )
    return '\n\n'.join(statements)


book = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = book.active
rows = list(sheet.iter_rows(min_row=4, values_only=True))
headers = list(rows[0])
index = {header: position for position, header in enumerate(headers)}
required_headers = {'سبب الاستبعاد', 'ID', 'الاسم العربي', 'الاسم الإنجليزي', 'الهاتف', 'النوع', 'البريد الإلكتروني', 'تاريخ الميلاد', 'الرقم القومي', 'الملاحظات'}
missing = required_headers - set(index)
if missing:
    raise ValueError(f'Missing expected columns: {sorted(missing)!r}')

selected = []
source_duplicates = 0
raw_phones = set()
for source_row, source in enumerate(rows[1:], start=5):
    reason = text(source[index['سبب الاستبعاد']])
    if 'رقم الهاتف مكرر' in reason:
        source_duplicates += 1
        continue

    legacy_id = text(source[index['ID']])
    name = text(source[index['الاسم العربي']]) or text(source[index['الاسم الإنجليزي']])
    gender_value = text(source[index['النوع']]).casefold()
    gender = {'0': 'female', 'false': 'female', '1': 'male', 'true': 'male'}.get(gender_value)
    raw_phone = re.sub(r'\s+', '', text(source[index['الهاتف']]).translate(TRANS))
    phone = raw_phone if len(raw_phone) <= 20 else ''
    notes = text(source[index['الملاحظات']])
    if raw_phone and not phone:
        notes = (notes + '\n' if notes else '') + f'هاتف المصدر: {raw_phone}'
    if not legacy_id or not name or not gender:
        raise ValueError(f'Cannot safely import Excel row {source_row}: missing ID, name, or valid gender')
    if phone and phone in raw_phones:
        raise ValueError(f'Unexpected duplicate phone among selected rows at Excel row {source_row}: {phone!r}')
    raw_phones.add(phone)
    selected.append([
        sql_text(legacy_id), sql_text(name), sql_text(phone), sql_text(source[index['البريد الإلكتروني']]),
        sql_text(gender), sql_text(source[index['الرقم القومي']]), sql_text(date_text(source[index['تاريخ الميلاد']])),
        sql_text(notes),
    ])

if source_duplicates == 0:
    raise ValueError('No duplicate-phone rows were found; stopping rather than preparing an unexpected upload.')

sql = f'''-- Adds only the members retained from excluded_members_for_client.xlsx.
-- It does not delete or update existing members, subscriptions, receipts, or leads.
-- Excluded because of duplicate phone in the source file: {source_duplicates} rows.
-- Intended additional members: {len(selected)} rows.
SET NAMES utf8mb4;
SET @target_branch_id := 1;
SET @expected_database := {sql_text(EXPECTED_DB)};

CREATE TEMPORARY TABLE `_fit90_additional_member_import` (
  `legacy_member_id` VARCHAR(30) NOT NULL,
  `name` VARCHAR(200) NOT NULL,
  `phone` VARCHAR(20) NULL,
  `email` VARCHAR(150) NULL,
  `gender` ENUM('male', 'female') NOT NULL,
  `card_number` VARCHAR(30) NULL,
  `date_of_birth` VARCHAR(10) NULL,
  `notes` TEXT NULL,
  PRIMARY KEY (`legacy_member_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

{insert_rows(selected)}

SELECT (BINARY DATABASE() = BINARY @expected_database) INTO @database_ok;
SELECT EXISTS(
  SELECT 1 FROM `tbl_branches`
  WHERE `branch_id` = @target_branch_id
    AND BINARY LOWER(TRIM(`branch_name`)) = BINARY 'fit90'
) INTO @branch_ok;
SET @is_safe := @database_ok AND @branch_ok;
SELECT @is_safe AS `precheck_passed`, @database_ok AS `database_ok`, @branch_ok AS `branch_ok`;

START TRANSACTION;

-- An already-imported member code or phone is skipped. Existing data is never overwritten.
INSERT INTO `club_members` (
  `member_code`, `name`, `phone`, `email`, `gender`, `card_number`, `date_of_birth`, `notes`,
  `branch_id`, `is_active`, `is_deleted`, `created_at`, `updated_at`
)
SELECT
  CONCAT('MIG-', `i`.`legacy_member_id`), `i`.`name`, `i`.`phone`, `i`.`email`, `i`.`gender`,
  `i`.`card_number`, `i`.`date_of_birth`, `i`.`notes`, @target_branch_id, 1, 0, NOW(3), NOW(3)
FROM `_fit90_additional_member_import` `i`
LEFT JOIN `club_members` `by_code`
  ON BINARY `by_code`.`member_code` = BINARY CONCAT('MIG-', `i`.`legacy_member_id`)
LEFT JOIN `club_members` `by_phone`
  ON `i`.`phone` IS NOT NULL AND BINARY `by_phone`.`phone` = BINARY `i`.`phone`
WHERE @is_safe = 1 AND `by_code`.`id` IS NULL AND `by_phone`.`id` IS NULL;

SELECT ROW_COUNT() AS `additional_members_inserted`;
{('ROLLBACK;' if DRY_RUN else 'COMMIT;')}

SELECT
  {len(selected)} AS `source_rows_requested`,
  {source_duplicates} AS `duplicate_phone_rows_excluded`,
  (SELECT COUNT(*) FROM `club_members` WHERE `branch_id` = @target_branch_id AND `member_code` LIKE 'MIG-%') AS `total_migrated_members_after_run`;

DROP TEMPORARY TABLE `_fit90_additional_member_import`;
'''

with open(OUTPUT, 'w', encoding='utf-8', newline='\n') as file:
    file.write(sql)

print({'output': OUTPUT, 'members_to_add': len(selected), 'duplicate_phone_rows_excluded': source_duplicates, 'dry_run': DRY_RUN})
