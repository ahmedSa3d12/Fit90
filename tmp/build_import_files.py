import csv, datetime, os, re, unicodedata
from openpyxl import load_workbook

OUT = r"E:\final_projects\asmaa\23-8-2026\FIT90_19-8\outputs\data-migration\upload-package-20260908"
MEMBERS = r"C:\Users\qqqqq\Downloads\Members-07-09-2026-01-17-42-PM.xlsx"
EXPORT = r"C:\Users\qqqqq\Downloads\Fit90_Full_Export_2026-09-07.xlsx"

def text(v):
    if v is None: return ''
    return unicodedata.normalize('NFKC', str(v)).replace('\u200b','').replace('\ufeff','').strip()
def digits(v):
    return text(v).translate(str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹','01234567890123456789'))
def phone(v):
    s=re.sub(r'[\s\-()+]','',digits(v))
    if re.fullmatch(r'201\d{9}',s): s='0'+s[2:]
    return s if re.fullmatch(r'01\d{9}',s) else ''
def date(v):
    if isinstance(v,(datetime.date,datetime.datetime)): return v.strftime('%Y-%m-%d')
    return text(v)
def read(path,sheet):
    wb=load_workbook(path,read_only=True,data_only=False,keep_links=False); ws=wb[sheet]
    h=[text(c.value) for c in next(ws.iter_rows(max_row=1))]
    for n,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2): yield n,dict(zip(h,row))
def write(name,fields,rows):
    with open(os.path.join(OUT,name),'w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore'); w.writeheader(); w.writerows(rows)

os.makedirs(OUT,exist_ok=True); rejected=[]
member_rows=[]; member_ids=set()
for n,r in read(MEMBERS,'Data'):
    legacy=text(r['ID']); member_ids.add(legacy); name=text(r['NameAR']) or text(r['NameEng']); p=phone(r['PhoneNo'])
    status='ready' if name and p else 'needs_review'
    row={'legacy_member_id':legacy,'target_branch_id':'1','name_ar':text(r['NameAR']),'name_en':text(r['NameEng']),'name':name,'phone':p,'phone_original':text(r['PhoneNo']),'gender_original':text(r['Gender']),'birth_date':date(r['BirthDate']),'email':text(r['Email']),'address':text(r['AthleticAddress']),'national_id':digits(r['NationalId']),'job_title':text(r['JobTitle']),'source_name':text(r['SourceName']),'notes':text(r['Comment']),'legacy_code':text(r['Code']),'creation_date_original':date(r['CreationDate']),'import_status':status}
    member_rows.append(row)
    if status!='ready': rejected.append({'entity':'member','source_file':os.path.basename(MEMBERS),'sheet':'Data','row_number':n,'legacy_id':legacy,'reason':'missing name or valid Egyptian mobile'})

sub_rows=[]
for n,r in read(EXPORT,'Memberships'):
    mid=text(r['memberId']); contract=text(r['contractNo']); status='ready' if mid in member_ids and contract else 'needs_review'
    row={'legacy_subscription_id':text(r['id']),'legacy_member_id':mid,'target_branch_id':'1','contract_number':contract,'access_code':text(r['accessCode']),'package_name':text(r['packageName']),'package_category':text(r['packageCategory']),'package_type_id_original':text(r['packageTypeId']),'legacy_status':text(r['statusName']),'start_date_original':date(r['startDateAsString']),'end_date_original':date(r['expirationDateAsString']),'payment_date_original':date(r['paymentDateAsString']),'creation_date_original':date(r['creationDate']),'price':text(r['price']),'price_before_discount':text(r['priceBeforeDiscount']),'discount':text(r['discount']),'discount_by_amount':text(r['discountByAmount']),'is_discount_percentage':text(r['isDiscountPercentage']),'amount_paid':text(r['totalAmountPaid']),'remaining_amount':text(r['totalAmountRemaining']),'receipt_number':text(r['receiptNo']),'payment_method_note':'payment method not present in source','sales_name_original':text(r['salesName']),'coach_name_original':text(r['coachName']),'attendance_count_original':text(r['attendanceCount']),'freeze_days_original':text(r['totalFreezedDays']),'notes':text(r['notes']),'import_status':status}
    sub_rows.append(row)
    if status!='ready': rejected.append({'entity':'subscription','source_file':os.path.basename(EXPORT),'sheet':'Memberships','row_number':n,'legacy_id':text(r['id']),'reason':'missing contract or unresolved source member'})

lead_rows=[]
for n,r in read(EXPORT,'Potential Members'):
    legacy=text(r['ID']); name=text(r['NameAR']) or text(r['NameEng']); p=phone(r['PhoneNo']); status='ready' if name and p else 'needs_review'
    row={'legacy_lead_id':legacy,'target_branch_id':'1','name_ar':text(r['NameAR']),'name_en':text(r['NameEng']),'name':name,'phone':p,'phone_original':text(r['PhoneNo']),'email':text(r['Email']),'gender_original':text(r['Gender']),'source_name_original':text(r['SourceName']),'status':'new','follow_up_date_original':date(r['LastCallDate']),'notes':text(r['Comment']),'creation_date_original':date(r['CreationDate']),'import_status':status}
    lead_rows.append(row)
    if status!='ready': rejected.append({'entity':'lead','source_file':os.path.basename(EXPORT),'sheet':'Potential Members','row_number':n,'legacy_id':legacy,'reason':'missing name or valid Egyptian mobile'})

write('members_import.csv',list(member_rows[0]),member_rows); write('subscriptions_import.csv',list(sub_rows[0]),sub_rows); write('potential_members_import.csv',list(lead_rows[0]),lead_rows)
write('needs_review.csv',['entity','source_file','sheet','row_number','legacy_id','reason'],rejected)
with open(os.path.join(OUT,'README.txt'),'w',encoding='utf-8') as f: f.write('Target branch: 1 (Fit90)\nCSV is UTF-8 with BOM. Preserve date_original fields until server-side date and package mapping validation is completed. Do not upload to production without dry run.\n')
print({'output':OUT,'members':len(member_rows),'subscriptions':len(sub_rows),'potential_members':len(lead_rows),'needs_review':len(rejected)})
