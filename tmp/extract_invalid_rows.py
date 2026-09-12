import datetime
import json
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

MEMBERS = r"C:\Users\qqqqq\Downloads\Members-07-09-2026-01-17-42-PM.xlsx"
EXPORT = r"C:\Users\qqqqq\Downloads\Fit90_Full_Export_2026-09-07.xlsx"
OUTPUT = r"E:\final_projects\asmaa\23-8-2026\FIT90_19-8\tmp\invalid_rows.json"
TYPES = {
    '1 month':8,'1 month premium':9,'golden month':10,'2 months':11,'3 months':12,
    '3 months premium':13,'3 months summer challange':14,'6 months':15,'6 months premium':16,
    'annual membership':17,'annual family membership':18,'morning annual membership':19,
    'morning membership 6am - 1 pm':20,'pre sale':21,'annual monglish offer':22,
    'students offer':23,'students offer moharem bek':24,'fit90 program':25,'marines':26,
    '1 session':27,'4 sessions':28,'8 sessions':29,'12 sessions':30,'16 sessions':31,
    '20 sessions':32,'40 sessions':33,'head coach 8 sessions':34,'head coach12 sessions':35,
    'head coach 20 sessions':36,'head coach 40 sessions':37,'youth program 10 sessions':38,
    'youth program 20 sessions':39,'8 sessions youth program head coach':40,
    '20 sessions youth program head coach':41,'nutrition 1 session':42,
    'nutrition 4 sessions':43,'nutrition 8 sessions':44,'nutrition 12 sessions':45,
    '4 sessions classes':46,'8 sessions classes':47,'12 sessions classes':48,
    '1 month full classes':49,
}
TRANS = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹','01234567890123456789')

def text(value):
    return '' if value is None else unicodedata.normalize('NFKC', str(value)).replace('\u200b','').replace('\ufeff','').strip()

def normalized(value):
    return re.sub(r'\s+', ' ', text(value)).casefold()

def phone(value):
    value = re.sub(r'[\s\-()+]', '', text(value).translate(TRANS))
    value = '0' + value[2:] if re.fullmatch(r'201\d{9}', value) else value
    value = '0' + value if re.fullmatch(r'1\d{9}', value) else value
    return value if re.fullmatch(r'01\d{9}', value) else ''

def date(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime('%Y-%m-%d')
    value = text(value)
    for pattern in ('%d-%m-%Y','%d-%m-%Y %I:%M %p','%Y-%m-%dT%H:%M:%S.%fZ'):
        try:
            return datetime.datetime.strptime(value, pattern).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return ''

def money(value):
    try:
        result = Decimal(text(value) or '0')
        return result if result >= 0 else Decimal('0')
    except (InvalidOperation, ValueError):
        return Decimal('0')

def rows(path, sheet):
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    worksheet = workbook[sheet]
    headers = [text(cell.value) for cell in next(worksheet.iter_rows(max_row=1))]
    for source_row, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        raw = {header: text(value) for header, value in zip(headers, values)}
        yield headers, source_row, raw

def report_row(source_row, raw, reason):
    return {'Source row': source_row, 'Reason for exclusion': '; '.join(reason), **raw}

member_headers = None
members_invalid = []
valid_members = {}
accepted_phones = set()
for member_headers, source_row, raw in rows(MEMBERS, 'Data'):
    legacy_id = text(raw.get('ID'))
    name = text(raw.get('NameAR')) or text(raw.get('NameEng'))
    normalized_phone = phone(raw.get('PhoneNo'))
    gender = {'0':'female','false':'female','1':'male','true':'male'}.get(text(raw.get('Gender')).casefold())
    reasons = []
    if not name: reasons.append('الاسم مفقود')
    if not normalized_phone: reasons.append('رقم الهاتف مفقود أو غير صالح')
    if not gender: reasons.append('النوع غير صالح؛ المقبول 0=أنثى أو 1=ذكر')
    if normalized_phone and normalized_phone in accepted_phones:
        reasons.append('رقم الهاتف مكرر؛ تم الاحتفاظ بأول عضو صحيح فقط')
    if reasons:
        members_invalid.append(report_row(source_row, raw, reasons))
        continue
    accepted_phones.add(normalized_phone)
    valid_members[legacy_id] = True

subscription_headers = None
subscriptions_invalid = []
for subscription_headers, source_row, raw in rows(EXPORT, 'Memberships'):
    member_id = text(raw.get('memberId'))
    package = text(raw.get('packageName'))
    start, end = date(raw.get('startDateAsString')), date(raw.get('expirationDateAsString'))
    contract = text(raw.get('contractNo'))
    reasons = []
    if member_id not in valid_members: reasons.append('العضو غير موجود ضمن الأعضاء المقبولين للاستيراد')
    if normalized(package) not in TYPES: reasons.append('نوع الاشتراك غير مطابق لأنواع الاشتراكات الموجودة بالنظام')
    if not start or not end or start > end: reasons.append('تاريخ بداية أو انتهاء الاشتراك غير صالح')
    if not contract: reasons.append('رقم العقد مفقود')
    if money(raw.get('totalAmountPaid')) > money(raw.get('price')) - money(raw.get('discountByAmount')):
        reasons.append('المبلغ المدفوع أكبر من قيمة الاشتراك بعد الخصم')
    if reasons:
        subscriptions_invalid.append(report_row(source_row, raw, reasons))

lead_headers = None
leads_invalid = []
for lead_headers, source_row, raw in rows(EXPORT, 'Potential Members'):
    name = text(raw.get('NameAR')) or text(raw.get('NameEng'))
    if not name:
        leads_invalid.append(report_row(source_row, raw, ['الاسم مفقود']))

result = {
    'summary': [
        {'File / sheet': 'Members-07-09-2026-01-17-42-PM.xlsx / Data', 'Invalid rows': len(members_invalid), 'Reason summary': 'اسم أو هاتف أو نوع مفقود/غير صالح، أو رقم هاتف مكرر'},
        {'File / sheet': 'Fit90_Full_Export_2026-09-07.xlsx / Memberships', 'Invalid rows': len(subscriptions_invalid), 'Reason summary': 'عضو غير مقبول، أو نوع اشتراك غير مطابق، أو تاريخ/عقد/مبلغ غير صالح'},
        {'File / sheet': 'Fit90_Full_Export_2026-09-07.xlsx / Potential Members', 'Invalid rows': len(leads_invalid), 'Reason summary': 'الاسم مفقود'},
    ],
    'members': members_invalid,
    'subscriptions': subscriptions_invalid,
    'leads': leads_invalid,
}
with open(OUTPUT, 'w', encoding='utf-8') as file:
    json.dump(result, file, ensure_ascii=False, indent=2)
print(json.dumps({key: len(result[key]) for key in ('members','subscriptions','leads')}, ensure_ascii=False))
