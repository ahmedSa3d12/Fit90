import csv, datetime, os, re, unicodedata
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

OUT=r'E:\final_projects\asmaa\23-8-2026\FIT90_19-8\outputs\data-migration\final-upload-package-20260908'; MEMBERS=r'C:\Users\qqqqq\Downloads\Members-07-09-2026-01-17-42-PM.xlsx'; EXPORT=r'C:\Users\qqqqq\Downloads\Fit90_Full_Export_2026-09-07.xlsx'
TYPES={'1 month':8,'1 month premium':9,'golden month':10,'2 months':11,'3 months':12,'3 months premium':13,'3 months summer challange':14,'6 months':15,'6 months premium':16,'annual membership':17,'annual family membership':18,'morning annual membership':19,'morning membership 6am - 1 pm':20,'pre sale':21,'annual monglish offer':22,'students offer':23,'students offer moharem bek':24,'fit90 program':25,'marines':26,'1 session':27,'4 sessions':28,'8 sessions':29,'12 sessions':30,'16 sessions':31,'20 sessions':32,'40 sessions':33,'head coach 8 sessions':34,'head coach12 sessions':35,'head coach 20 sessions':36,'head coach 40 sessions':37,'youth program 10 sessions':38,'youth program 20 sessions':39,'8 sessions youth program head coach':40,'20 sessions youth program head coach':41,'nutrition 1 session':42,'nutrition 4 sessions':43,'nutrition 8 sessions':44,'nutrition 12 sessions':45,'4 sessions classes':46,'8 sessions classes':47,'12 sessions classes':48,'1 month full classes':49}
TRANS=str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹','01234567890123456789')
def t(x): return '' if x is None else unicodedata.normalize('NFKC',str(x)).replace('\u200b','').replace('\ufeff','').strip()
def norm(x): return re.sub(r'\s+',' ',t(x)).casefold()
def ph(x):
 s=re.sub(r'[\s\-()+]','',t(x).translate(TRANS)); s='0'+s[2:] if re.fullmatch(r'201\d{9}',s) else s; s='0'+s if re.fullmatch(r'1\d{9}',s) else s
 return s if re.fullmatch(r'01\d{9}',s) else ''
def d(x):
 if isinstance(x,(datetime.date,datetime.datetime)): return x.strftime('%Y-%m-%d')
 s=t(x)
 for f in ('%d-%m-%Y','%d-%m-%Y %I:%M %p','%Y-%m-%dT%H:%M:%S.%fZ'):
  try:return datetime.datetime.strptime(s,f).strftime('%Y-%m-%d')
  except ValueError:pass
 return ''
def money(x):
 try:
  value=Decimal(t(x) or '0')
  return value if value >= 0 else Decimal('0')
 except (InvalidOperation, ValueError): return Decimal('0')
def money_text(value):
 text=format(value, 'f')
 return (text.rstrip('0').rstrip('.') if '.' in text else text) or '0'
def read(path,sheet):
 w=load_workbook(path,read_only=True,data_only=False,keep_links=False); s=w[sheet]; h=[t(c.value) for c in next(s.iter_rows(max_row=1))]
 for n,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):yield n,dict(zip(h,row))
def write(name,rows):
 with open(os.path.join(OUT,name),'w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]) if rows else []);w.writeheader();w.writerows(rows)
os.makedirs(OUT,exist_ok=True); review=[]; members=[]; valid={}; accepted_phones=set()
for n,r in read(MEMBERS,'Data'):
 i=t(r['ID']); name=t(r['NameAR']) or t(r['NameEng']); p=ph(r['PhoneNo']); gender={'0':'female','false':'female','1':'male','true':'male'}.get(t(r['Gender']).casefold())
 if not(name and p and gender):review.append({'entity':'member','source_row':n,'legacy_id':i,'reason':'missing required name, phone, or gender'});continue
 if p in accepted_phones:review.append({'entity':'member','source_row':n,'legacy_id':i,'reason':'duplicate phone; first valid member retained'});continue
 accepted_phones.add(p)
 valid[i]=True;members.append({'legacy_member_id':i,'branch_id':1,'name':name,'phone':p,'gender':gender,'date_of_birth':d(r['BirthDate']),'email':t(r['Email']),'address':t(r['AthleticAddress']),'national_id':t(r['NationalId']).translate(TRANS),'job_title':t(r['JobTitle']),'notes':t(r['Comment'])})
subs=[]; used=set()
for n,r in read(EXPORT,'Memberships'):
 i=t(r['id']);mid=t(r['memberId']);package_id=TYPES.get(norm(r['packageName']));start,end=d(r['startDateAsString']),d(r['expirationDateAsString']);contract=t(r['contractNo']); reasons=[]
 if mid not in valid:reasons.append('member not accepted')
 if not package_id:reasons.append('package not matched locally')
 if not(start and end and start<=end):reasons.append('invalid subscription dates')
 if not contract:reasons.append('missing contract')
 value,discount,paid=money(r['price']),money(r['discountByAmount']),money(r['totalAmountPaid'])
 if paid > value-discount:reasons.append('paid amount exceeds net value')
 if reasons:review.append({'entity':'subscription','source_row':n,'legacy_id':i,'reason':'; '.join(reasons)});continue
 number=contract if contract not in used and len(contract)<=30 else f'LEG-{i}';used.add(number)
 subs.append({'legacy_subscription_id':i,'legacy_member_id':mid,'branch_id':1,'subscription_number':number,'subscription_type_id':package_id,'subscription_type_name':t(r['packageName']),'registration_date':d(r['paymentDateAsString']) or d(r['creationDate']),'subscription_start_date':start,'subscription_end_date':end,'subscription_value':money_text(value),'discount_value':money_text(discount),'paid_amount':money_text(paid),'remaining_amount':money_text(value-discount-paid),'legacy_remaining_amount':t(r['totalAmountRemaining']),'legacy_receipt_number':t(r['receiptNo']),'legacy_status':t(r['statusName']),'notes':t(r['notes'])})
leads=[]
for n,r in read(EXPORT,'Potential Members'):
 i=t(r['ID']);name=t(r['NameAR']) or t(r['NameEng'])
 if not name:review.append({'entity':'lead','source_row':n,'legacy_id':i,'reason':'missing name'});continue
 leads.append({'legacy_lead_id':i,'branch_id':1,'name':name,'phone':ph(r['PhoneNo']),'email':t(r['Email']),'gender':{'0':'female','false':'female','1':'male','true':'male'}.get(t(r['Gender']).casefold(),'') ,'source_name':t(r['SourceName']),'status':'new','follow_up_at':d(r['LastCallDate']),'notes':t(r['Comment'])})
write('members_final.csv',members);write('subscriptions_final.csv',subs);write('potential_members_final.csv',leads);write('needs_review.csv',review)
open(os.path.join(OUT,'README.txt'),'w',encoding='utf8').write('Validated against local branch 1. Import members first, persist legacy_member_id to new ID crosswalk, then subscriptions, then leads. Do not import directly with phpMyAdmin. remaining_amount is recalculated as subscription_value - discount_value - paid_amount; legacy_remaining_amount preserves the original exported value.\n')
print({'members':len(members),'subscriptions':len(subs),'leads':len(leads),'review':len(review)})
