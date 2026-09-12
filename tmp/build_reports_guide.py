from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(r"E:\final_projects\asmaa\FIT90\documentation\FIT90-دليل-إدارة-التقارير.docx")

NAVY = "1E3A5F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "E8F5E9"
PALE_GOLD = "FFF6DC"
MUTED = "666666"
WHITE = "FFFFFF"
BLACK = "111111"


reports = [
    {
        "category": "التقارير المالية والدخل",
        "name": "الأرباح",
        "path": "/mos/reports/profit",
        "purpose": "ملخص الربحية التشغيلية خلال الفترة.",
        "source": "إيصالات الاشتراكات + القيود المالية من نوع إيراد آخر + القيود من نوع مصروف.",
        "conditions": "يستبعد القيود المحذوفة، ويطبّق فترة الإيصال على الإيصالات وفترة القيد على الإيرادات الأخرى والمصروفات.",
        "formula": "إجمالي الدخل = مجموع الإيصالات + مجموع الإيرادات الأخرى. صافي الربح = إجمالي الدخل − مجموع المصروفات.",
        "filters": "من تاريخ، إلى تاريخ، والفرع عند تمريره.",
        "note": "نطاق المصروف هنا هو entry_type = expense فقط؛ الرواتب والعمولات والمكافآت لا تُخصم إلا إذا سُجلت أيضًا كمصروف.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "المصروفات",
        "path": "/mos/reports/expenses",
        "purpose": "عرض المصروفات المسجلة مع إجماليها.",
        "source": "club_financial_entries.",
        "conditions": "entry_type = expense و is_deleted = false.",
        "formula": "الإجمالي = مجموع amount لكل المصروفات المطابقة، وليس مجموع الصفحة الحالية فقط. العدد = عدد القيود المطابقة.",
        "filters": "من تاريخ، إلى تاريخ، والفرع عند تمريره.",
        "note": "زر «إضافة مصروف» أعلى التقرير ينقل إلى /mos/accounts/expenses للتسجيل والتعديل والحذف.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "دخل اشتراكات الجيم",
        "path": "/mos/reports/membershipsIncome",
        "purpose": "حصر المقبوضات الخاصة بعضويات الجيم العادية.",
        "source": "إيصالات مرتبطة باشتراكات وبنوع باقة.",
        "conditions": "الاشتراك غير مرتبط بجلسات، وتصنيف الباقة ليس private أو medical.",
        "formula": "الإجمالي = مجموع مبالغ الإيصالات المطابقة. العدد = عدد الإيصالات.",
        "filters": "تاريخ الإيصال والفرع.",
        "note": "يفصل دخل الجيم عن دخل التدريب الشخصي والعيادة لمنع التداخل.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "دخل التدريب الشخصي",
        "path": "/mos/reports/privateMembershipsIncome",
        "purpose": "إجمالي المقبوضات الخاصة بالتدريب الشخصي.",
        "source": "إيصالات الاشتراكات.",
        "conditions": "تصنيف الباقة private أو الاشتراك مرتبط بالجلسات، لدعم البيانات القديمة والجديدة.",
        "formula": "يجمع الإيصالات لكل اشتراك؛ مبلغ الصف = مجموع إيصالات الاشتراك، والعدد = عدد إيصالاته. ملخص التقرير يجمع كل الإيصالات.",
        "filters": "تاريخ الإيصال والفرع.",
        "note": "التجميع في الجدول على مستوى الاشتراك وليس على مستوى الإيصال.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "دخل اشتراكات العيادة",
        "path": "/mos/reports/medicalMembershipsIncome",
        "purpose": "عرض دخل الباقات الطبية فقط.",
        "source": "إيصالات الاشتراكات المرتبطة بأنواع الباقات.",
        "conditions": "package_category = medical بصرف النظر عن اسم الباقة بالعربية أو الإنجليزية.",
        "formula": "الإجمالي = مجموع مبالغ الإيصالات الطبية. العدد = عدد الإيصالات.",
        "filters": "تاريخ الإيصال والفرع.",
        "note": "كان التقرير سابقًا يبحث عن كلمة «طبي/medical» داخل الاسم؛ تم تصحيحه ليستخدم تصنيف الباقة.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "الأرباح يومًا بيوم",
        "path": "/mos/reports/daybydayprofit",
        "purpose": "تحليل الربح التشغيلي لكل يوم.",
        "source": "الإيصالات، الإيرادات الأخرى، والمصروفات.",
        "conditions": "القيود غير المحذوفة، مع تجميع مستقل حسب تاريخ كل يوم.",
        "formula": "دخل اليوم = إيصالات اليوم + الإيرادات الأخرى. صافي اليوم = دخل اليوم − مصروفات اليوم. الملخص يجمع الفترة كاملة.",
        "filters": "من تاريخ وإلى تاريخ؛ الافتراضي من سبعة أيام سابقة حتى اليوم.",
        "note": "تم تصحيحه؛ لم يعد يعرض الإيصالات وحدها تحت مسمى الربح.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "التقرير المالي للموظفين",
        "path": "/mos/reports/employeeFinancial",
        "purpose": "كل القيود المالية المرتبطة بموظف.",
        "source": "club_financial_entries.",
        "conditions": "employee_id غير فارغ و is_deleted = false.",
        "formula": "الإجمالي = مجموع مبالغ جميع القيود المطابقة، والعدد = عددها.",
        "filters": "تاريخ القيد والفرع.",
        "note": "يشمل الأنواع المختلفة للقيود المرتبطة بالموظف، وليس الرواتب فقط.",
    },
    {
        "category": "التقارير المالية والدخل",
        "name": "دخل الاشتراكات لكل نوع عضوية",
        "path": "/mos/reports/membershipsIncomePerPackageType",
        "purpose": "مقارنة قيمة ومدفوعات الاشتراكات حسب اسم نوع العضوية.",
        "source": "سجلات الاشتراكات.",
        "conditions": "التجميع حسب subscription_type للاشتراكات المسجلة في الفترة.",
        "formula": "لكل نوع: العدد، مجموع paid_amount كدخل محصل، ومجموع subscription_value كقيمة تعاقدية. الإجمالي العام = مجموع المدفوع.",
        "filters": "تاريخ تسجيل الاشتراك والفرع.",
        "note": "يعتمد على رصيد paid_amount الحالي للاشتراك، وليس تاريخ كل دفعة منفردة.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "الاشتراكات",
        "path": "/mos/reports/memberships/all",
        "purpose": "القائمة الشاملة للاشتراكات وموقفها المالي.",
        "source": "club_subscriptions.",
        "conditions": "يعرض كل الاشتراكات المطابقة، ويحوّل الاشتراك active المنتهي تاريخيًا إلى expired للعرض.",
        "formula": "إجمالي القيمة = مجموع subscription_value، المدفوع = مجموع paid_amount، المتبقي = مجموع remaining_amount.",
        "filters": "تاريخ التسجيل، الفرع، والبحث بالاسم أو رقم الاشتراك أو نوعه.",
        "note": "الإجماليات محسوبة على كامل النتائج لا الصفحة الحالية.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "التدريب الشخصي",
        "path": "/mos/reports/memberships/privateMemberships",
        "purpose": "الاشتراكات الخاصة أو المرتبطة بالجلسات.",
        "source": "club_subscriptions مع نوع الباقة.",
        "conditions": "package_category = private أو is_linked_to_sessions = true.",
        "formula": "الإيراد = مجموع paid_amount، والقيمة = مجموع subscription_value، مع عرض الجلسات الممنوحة والمستخدمة.",
        "filters": "تاريخ التسجيل والفرع.",
        "note": "يدعم تصنيف الباقات الجديد والبيانات القديمة المرتبطة بالجلسات.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "الاشتراكات غير المجددة",
        "path": "/mos/reports/memberships/notRenewed",
        "purpose": "الأعضاء الذين انتهى اشتراكهم ولم يغطهم اشتراك أحدث ساري.",
        "source": "سجل الاشتراكات الكامل لكل عضو.",
        "conditions": "تاريخ الانتهاء قبل اليوم، ولا يوجد للعضو اشتراك ينتهي اليوم أو بعده. يحتفظ بأحدث اشتراك منتهي لكل عضو.",
        "formula": "العدد = عدد الأعضاء/الاشتراكات النهائية بعد حذف من لديهم تغطية سارية.",
        "filters": "فترة تاريخ الانتهاء والفرع.",
        "note": "الفترة تُطبّق على تاريخ انتهاء الاشتراك.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "الاشتراكات الجديدة والمجددة",
        "path": "/mos/reports/memberships/newRenewed",
        "purpose": "تقسيم اشتراكات الفترة إلى جديدة وتجديد.",
        "source": "اشتراكات الفترة مع التاريخ الكامل لاشتراكات نفس الأعضاء.",
        "conditions": "الاشتراك «مجدد» إذا وُجد للعضو اشتراك أقدم في التاريخ أو برقم سجل أقل في اليوم نفسه؛ وإلا فهو «جديد».",
        "formula": "new = عدد الاشتراكات بلا سابقة، renewed = عدد الاشتراكات ذات السابقة، total = مجموعهما.",
        "filters": "تاريخ التسجيل والفرع.",
        "note": "تم ضبط ترقيم الصفحات بعد التصنيف.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "سجل الاشتراكات",
        "path": "/mos/reports/memberships/log",
        "purpose": "سجل زمني لإنشاء الاشتراكات.",
        "source": "club_subscriptions.",
        "conditions": "الترتيب حسب تاريخ التسجيل ثم رقم السجل تنازليًا.",
        "formula": "العدد = عدد الاشتراكات المطابقة، مع عرض القيمة والحالة المحسوبة.",
        "filters": "تاريخ التسجيل والفرع.",
        "note": "هذا سجل الاشتراكات، وليس سجل تدقيق لكل تعديل حدث عليها.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "استخدام العضويات",
        "path": "/mos/reports/packagesUntil",
        "purpose": "تقسيم الاشتراكات السارية حسب قرب انتهاء الصلاحية.",
        "source": "الاشتراكات التي ينتهي تاريخها اليوم أو بعده.",
        "conditions": "ثلاث فئات: خلال 7 أيام، من اليوم الثامن حتى 30 يومًا، وبعد 30 يومًا.",
        "formula": "لكل فئة: عدد الاشتراكات ومجموع subscription_value.",
        "filters": "الفرع فقط؛ التقرير نافذة مستقبلية ثابتة لذلك لا يظهر فلتر التاريخ.",
        "note": "القيمة المعروضة قيمة الاشتراكات وليست المقبوضات.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "نقل الاشتراكات",
        "path": "/club/subscriptions/transfers",
        "purpose": "شاشة تشغيلية لعرض وإدارة عمليات نقل الاشتراك.",
        "source": "club_subscription_transfers.",
        "conditions": "ليست report handler عامًا؛ هي شاشة متخصصة ضمن قسم التقارير.",
        "formula": "تعرض بيانات النقل من الباقة القديمة إلى الجديدة والقيم والتاريخ، بلا معادلة إجمالي موحدة.",
        "filters": "فلاتر شاشة النقل نفسها.",
        "note": "وجودها بالقسم لتسهيل الوصول، لكنها شاشة عمليات وليست تقريرًا تجميعيًا.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "تطوير الاشتراكات",
        "path": "/mos/reports/membershipUpgrade",
        "purpose": "استخراج عمليات النقل التي تمثل ترقية مالية.",
        "source": "club_subscription_transfers.",
        "conditions": "to_value أكبر من from_value.",
        "formula": "قيمة الترقية = to_value − from_value. الإجمالي = مجموع فروق كل الترقيات.",
        "filters": "تاريخ النقل والفرع.",
        "note": "النقل المتساوي أو الأقل قيمة لا يُعد ترقية.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "الخصومات على الاشتراكات",
        "path": "/mos/reports/membershipsDiscount",
        "purpose": "الاشتراكات التي طُبق عليها خصم.",
        "source": "club_subscriptions.",
        "conditions": "discount_enabled = true و discount_value > 0.",
        "formula": "إجمالي الخصم = مجموع discount_value لكل النتائج.",
        "filters": "تاريخ التسجيل والفرع.",
        "note": "يعرض قيمة الاشتراك والمدفوع والخصم للمراجعة.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "أقصى تاريخ انتهاء",
        "path": "/mos/reports/maximumExpirationDate",
        "purpose": "آخر تاريخ انتهاء معروف لكل عضو.",
        "source": "كل اشتراكات الأعضاء.",
        "conditions": "يحتفظ بالاشتراك صاحب أكبر subscription_end_date لكل member_id.",
        "formula": "صف واحد لكل عضو، ويعرض هل التاريخ ما زال ساريًا مقارنة باليوم.",
        "filters": "فترة أقصى تاريخ انتهاء والفرع.",
        "note": "الفترة تُطبّق بعد تحديد أقصى تاريخ فعلي لكل عضو حتى لا تختار اشتراكًا أقدم خطأً.",
    },
    {
        "category": "تقارير الاشتراكات",
        "name": "استهلاك المزايا",
        "path": "/mos/reports/benefitsConsumption",
        "purpose": "مقارنة المزايا الممنوحة بالمستهلك والمتبقي لكل اشتراك.",
        "source": "مزايا الاشتراك وسجلات المزايا والحجوزات المكتملة والتجميد والدعوات والجلسات.",
        "conditions": "الاشتراك يحتوي على ميزة ممنوحة أكبر من صفر، ويُقارن الاستهلاك داخل مدة الاشتراك فقط.",
        "formula": "المتبقي = الممنوح − المستخدم. لتجنب التكرار يأخذ أقوى قيمة موثقة من مصادر الاستهلاك المتوازية بدل جمع نفس العملية مرتين.",
        "filters": "تداخل مدة الاشتراك مع الفترة، الفرع، والبحث بالعضو أو الكود أو الاشتراك أو النوع.",
        "note": "المستخدم لا يتجاوز الممنوح، والمتبقي لا يقل عن صفر.",
    },
    {
        "category": "تقارير الحضور والأعضاء",
        "name": "أعياد الميلاد",
        "path": "/mos/reports/birthdays",
        "purpose": "قائمة الأعضاء الذين لديهم تاريخ ميلاد مسجل.",
        "source": "club_members.",
        "conditions": "is_deleted = false و date_of_birth غير فارغ.",
        "formula": "العدد = عدد الأعضاء المطابقين.",
        "filters": "الفرع فقط؛ لا يظهر فلتر التاريخ لأن سنة الميلاد التاريخية لا تساوي سنة التقرير.",
        "note": "مرتب حسب تاريخ الميلاد المسجل.",
    },
    {
        "category": "تقارير الحضور والأعضاء",
        "name": "الأعضاء غير النشطين",
        "path": "/mos/reports/inactiveMembers",
        "purpose": "الأعضاء الذين لم يسجلوا حضورًا خلال فترة المتابعة.",
        "source": "الأعضاء وعلاقة الحضور.",
        "conditions": "العضو غير محذوف ولا توجد له أي حركة حضور بين بداية ونهاية الفترة.",
        "formula": "العدد = count مباشر من قاعدة البيانات للأعضاء بلا حضور، قبل تطبيق ترقيم الصفحات.",
        "filters": "من تاريخ وإلى تاريخ والفرع؛ الافتراضي آخر 30 يومًا.",
        "note": "تم تصحيح الفلترة لتحدث في قاعدة البيانات بدل فلترة الصفحة بعد جلبها.",
    },
    {
        "category": "تقارير الحضور والأعضاء",
        "name": "الحضور الزائد",
        "path": "/mos/reports/overAttendance",
        "purpose": "كشف العضو الذي سجل أكثر من حضور في اليوم نفسه.",
        "source": "club_attendance.",
        "conditions": "التجميع حسب member_id + attendance_date والاحتفاظ بالعدد الأكبر من 1.",
        "formula": "checkIns = عدد حركات العضو في اليوم. count = عدد حالات العضو/اليوم المخالفة.",
        "filters": "تاريخ الحضور والفرع.",
        "note": "يعرض الاسم والتاريخ وعدد مرات الدخول.",
    },
    {
        "category": "تقارير الحضور والأعضاء",
        "name": "الأعضاء الأكثر حضورًا",
        "path": "/mos/reports/topActiveMembers",
        "purpose": "ترتيب الأعضاء حسب عدد مرات الحضور.",
        "source": "club_attendance.",
        "conditions": "تجميع حسب member_id ثم ترتيب تنازلي، وبحد أقصى أعلى 50 عضوًا.",
        "formula": "checkIns لكل عضو = عدد حركاته. ملخص التقرير يعرض عدد الأعضاء وإجمالي الحركات.",
        "filters": "تاريخ الحضور والفرع.",
        "note": "ترقيم الصفحات يعمل داخل قائمة أعلى 50.",
    },
    {
        "category": "تقارير الحضور والأعضاء",
        "name": "حضور الجيم",
        "path": "/mos/reports/gymAttendanceCount",
        "purpose": "رقم إجمالي سريع لحركات دخول الجيم.",
        "source": "club_attendance.",
        "conditions": "كل حركة حضور مطابقة للفترة والفرع.",
        "formula": "count = عدد سجلات الحضور، وليس عدد الأعضاء الفريدين.",
        "filters": "تاريخ الحضور والفرع.",
        "note": "دخول العضو مرتين يُحسب حركتين.",
    },
    {
        "category": "تقارير الحضور والأعضاء",
        "name": "الحضور المتكرر في اليوم",
        "path": "/mos/reports/multipleAttendancePerDay",
        "purpose": "تفاصيل حالات تعدد الحضور في اليوم مع كود العضو.",
        "source": "club_attendance.",
        "conditions": "التجميع حسب اليوم والعضو، وإظهار الحالات ذات attendanceCount > 1.",
        "formula": "totalIncidents = عدد تركيبات العضو/اليوم التي تكرر فيها الحضور.",
        "filters": "تاريخ الحضور والفرع.",
        "note": "يشبه تقرير الحضور الزائد لكنه يضيف كود العضو ويستخدم تسمية أوضح للحالة.",
    },
    {
        "category": "تقارير المبيعات والعمولات",
        "name": "عمولة موظفي المبيعات",
        "path": "/mos/reports/salesCommission",
        "purpose": "حساب عمولة مسؤولي التسويق/المبيعات النشطين.",
        "source": "الموظفون، المسميات الوظيفية، والاشتراكات المنسوبة إلى sales_id.",
        "conditions": "موظف نشط من نوع الموظفين ومسماه مطابق لمسؤول التسويق، والاشتراك مسجل في الفترة.",
        "formula": "العمولة = المدفوع المحصل × نسبة عمولة الموظف ÷ 100. تحقيق المستهدف = المدفوع ÷ المستهدف × 100.",
        "filters": "تاريخ تسجيل الاشتراك والفرع.",
        "note": "إجمالي المبيعات تعاقدي، أما العمولة فتُحسب على المدفوع المحصل.",
    },
    {
        "category": "تقارير المبيعات والعمولات",
        "name": "عمولة العضويات",
        "path": "/mos/reports/customPackagesCommission",
        "purpose": "نسبة عمولة المدرب على اشتراكات الأعضاء المسندين إليه.",
        "source": "موظفو التدريب، ملفات المدربين، الأعضاء، والاشتراكات.",
        "conditions": "الموظف مدرب نشط، والعضو مرتبط بملف المدرب، والاشتراك ضمن الفترة.",
        "formula": "عمولة المدرب = مدفوع اشتراكات أعضائه × نسبة عمولته ÷ 100. تحقيق المستهدف بنفس قاعدة المدفوع.",
        "filters": "تاريخ تسجيل الاشتراك والفرع.",
        "note": "النسبة والمستهدف مأخوذان من ملف موظف المدرب.",
    },
    {
        "category": "تقارير المبيعات والعمولات",
        "name": "استخدام الفترات لموظفي المبيعات",
        "path": "/mos/reports/packageUtilizationPerSalesPersonal",
        "purpose": "توزيع أنواع الباقات المباعة على موظفي المبيعات.",
        "source": "الاشتراكات وموظفو المبيعات.",
        "conditions": "التجميع حسب موظف المبيعات + subscription_type، مع فئة «غير مسند» عند غياب الموظف.",
        "formula": "count = عدد الاشتراكات، total = مجموع subscription_value لكل مجموعة.",
        "filters": "تاريخ تسجيل الاشتراك والفرع.",
        "note": "total قيمة تعاقدية وليست المدفوع المحصل.",
    },
    {
        "category": "تقارير المبيعات والعمولات",
        "name": "نسبة إقناع موظف المبيعات",
        "path": "/mos/reports/salesPersonClosingRatio",
        "purpose": "قياس تحويل العملاء المحتملين لكل مسؤول مبيعات.",
        "source": "club_leads و club_sales_staff.",
        "conditions": "الحالات المحولة: converted أو won أو subscribed أو closed أو closed_won.",
        "formula": "نسبة الموظف = عدد العملاء المحولين ÷ إجمالي العملاء المسندين × 100. النسبة العامة بنفس المعادلة على الإجمالي.",
        "filters": "تاريخ إنشاء العميل المحتمل والفرع.",
        "note": "تم تصحيح فلتر حقول DateTime ليغطي اليوم كاملًا دون خطأ تشغيل.",
    },
    {
        "category": "تقارير المبيعات والعمولات",
        "name": "تفاصيل نسبة إقناع موظف المبيعات",
        "path": "/mos/reports/salesPersonClosingRatioDetails",
        "purpose": "عرض كل عميل محتمل وحالته وعدد المكالمات.",
        "source": "العملاء المحتملون، موظفو المبيعات، والمكالمات.",
        "conditions": "يربط المكالمات بالعميل عن طريق رقم الهاتف لعدم وجود lead_id مباشر.",
        "formula": "callsCount = عدد المكالمات المطابقة لنفس الهاتف داخل الفترة.",
        "filters": "تاريخ إنشاء العميل والمكالمات والفرع.",
        "note": "المطابقة بالهاتف تقريبية وتتطلب توحيد صيغة أرقام الهواتف.",
    },
    {
        "category": "تقارير المبيعات والعمولات",
        "name": "السجلات",
        "path": "/mos/reports/logs",
        "purpose": "سجل النشاط التشغيلي.",
        "source": "business_audit_log أولًا؛ وعند عدم وجود سجلات يستخدم آخر الحضور والإيصالات كبديل.",
        "conditions": "سجل التدقيق حسب الفترة والفرع. البديل محدود بآخر 200 حضور و200 إيصال.",
        "formula": "العدد = عدد سجلات التدقيق المطابقة، أو حجم موجز النشاط البديل بعد الدمج والترتيب.",
        "filters": "تاريخ الإنشاء/الحضور والفرع.",
        "note": "تم تصحيح فلتر التاريخ وترقيم الصفحات؛ سجل التدقيق هو المصدر المفضل.",
    },
]


def set_font(run, size=11, bold=False, color=BLACK, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    set_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run("صفحة ")
    set_font(run, 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_para(doc, text="", size=11, bold=False, color=BLACK, after=6, before=0,
             align=WD_ALIGN_PARAGRAPH.RIGHT, keep=False):
    p = doc.add_paragraph()
    set_rtl(p, align)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_labeled(doc, label, text):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(f"{label}: ")
    set_font(r1, size=10.5, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_font(r2, size=10.5, color=BLACK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(p)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}[level], bold=True,
             color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])
    return p


def add_status_box(doc, text, fill=PALE_GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    set_rtl(p)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=10.5, bold=True, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
]:
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
set_rtl(hp)
hp.paragraph_format.space_after = Pt(0)
hr = hp.add_run("FIT90  |  دليل إدارة التقارير")
set_font(hr, size=9, bold=True, color=MUTED)
footer = section.footer
add_page_number(footer.paragraphs[0])

# Editorial cover
add_para(doc, "FIT90", size=12, bold=True, color=BLUE, after=60, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "دليل إدارة التقارير", size=28, bold=True, color=NAVY, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "مصادر البيانات، شروط الظهور، وطريقة حساب كل تقرير",
         size=14, color=DARK_BLUE, after=28, align=WD_ALIGN_PARAGRAPH.CENTER)
add_status_box(doc, "نتيجة المراجعة: تم تشغيل تقارير القسم على قاعدة البيانات المحلية، وإصلاح أخطاء الفلاتر والحساب والترقيم.")
add_para(doc, "تاريخ المراجعة: 29 يوليو 2026", size=11, bold=True, color=MUTED,
         after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "النطاق: 31 عنصرًا ظاهرًا داخل «إدارة التقارير»", size=11, color=MUTED,
         after=70, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "إعداد فني لنظام FIT90", size=10, color=MUTED,
         align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()
add_heading(doc, "1. نتيجة المراجعة", 1)
add_status_box(doc, "الحالة النهائية: جميع عناصر قسم التقارير مرتبطة بمسار فعلي، وتقارير الـ API المختبرة تعمل دون أخطاء تشغيل ضمن فترة التاريخ.")
add_labeled(doc, "الاختبارات", "نجحت 12 حزمة اختبار بإجمالي 80 اختبارًا.")
add_labeled(doc, "البناء", "نجح بناء الـ Backend وبناء الواجهة Frontend.")
add_labeled(doc, "اختبار البيانات", "تم تشغيل تقارير القسم على قاعدة البيانات المحلية مرة دون فترة ومرة على الفترة 01–31 يوليو 2026.")
add_labeled(doc, "ترقيم الصفحات", "تم التحقق أن التقرير لا يعيد أكثر من حجم الصفحة المطلوب، وأن الإجماليات لا تعتمد على الصفحة الحالية.")
add_labeled(doc, "التنسيق", "تم تنسيق القيم الرقمية إلى منزلتين عشريتين كحد أقصى، بما في ذلك تسميات الرسوم البيانية.")

add_heading(doc, "2. أهم التصحيحات المنفذة", 1)
corrections = [
    ("فلتر التاريخ", "إصلاح حقول DateTime في نسبة إقناع المبيعات وتفاصيلها والسجلات لتغطية اليوم كاملًا."),
    ("الأرباح اليومية", "إضافة الإيرادات الأخرى وطرح المصروفات يوميًا بدل عرض الإيصالات وحدها."),
    ("تصنيف الدخل", "فصل دخل الجيم عن التدريب الشخصي والعيادة، وربط العيادة بتصنيف medical."),
    ("الأعضاء غير النشطين", "نقل شرط عدم الحضور إلى استعلام قاعدة البيانات للحصول على عدد وصفحات صحيحة."),
    ("الإجماليات", "حساب إجماليات المصروفات والديون من كامل النتائج."),
    ("الصفحات", "تصحيح تقارير الاشتراكات الجديدة، السجلات، الحضور المتكرر، الأكثر حضورًا، وتفاصيل المبيعات."),
    ("واجهة البحث", "إخفاء بحث الجدول من التقارير التي لا تنفذه حتى لا يعطي انطباعًا مضللًا."),
]
for label, text in corrections:
    add_labeled(doc, label, text)

add_heading(doc, "3. قواعد عامة لقراءة التقارير", 1)
add_labeled(doc, "الفترة", "كل تقرير يوضح أدناه الحقل الذي تُطبق عليه الفترة؛ قد يكون تاريخ الإيصال أو التسجيل أو الانتهاء أو الحضور أو القيد.")
add_labeled(doc, "الفرع", "عند تمرير branchId يقتصر الاستعلام على الفرع المحدد.")
add_labeled(doc, "الإجمالي", "بطاقات الملخص تحسب كامل النتائج المطابقة، بينما الجدول يعرض الصفحة الحالية.")
add_labeled(doc, "الحذف", "السجلات التي تحمل is_deleted = true تُستبعد حيث يدعم الجدول الحذف المنطقي.")
add_labeled(doc, "مبالغ الاشتراك", "subscription_value قيمة تعاقدية، paid_amount مبلغ محصل، remaining_amount متبقٍ على الاشتراك.")

doc.add_page_break()
add_heading(doc, "4. فهرس تقارير القسم", 1)
table = doc.add_table(rows=1, cols=3)
headers = ["م", "التقرير", "المسار"]
for idx, text in enumerate(headers):
    cell = table.rows[0].cells[idx]
    shade_cell(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    set_font(r, 10, True, NAVY)
set_repeat_table_header(table.rows[0])
for index, report in enumerate(reports, 1):
    cells = table.add_row().cells
    values = [str(index), report["name"], report["path"]]
    for col, value in enumerate(values):
        p = cells[col].paragraphs[0]
        set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(value)
        set_font(r, 9.2 if col == 2 else 9.5, color=BLACK)
        if index % 2 == 0:
            shade_cell(cells[col], LIGHT_GRAY)
set_table_geometry(table, [600, 3200, 5560])

current_category = None
report_number = 0
for report in reports:
    if report["category"] != current_category:
        doc.add_page_break()
        current_category = report["category"]
        add_heading(doc, current_category, 1)
        add_para(doc, "يوضح هذا القسم من أين تأتي البيانات وكيف تُحسب النتيجة وما الذي تعنيه الفلاتر.",
                 size=10.5, color=MUTED, after=12)
    report_number += 1
    add_heading(doc, f"{report_number}. {report['name']}", 2)
    add_labeled(doc, "المسار", report["path"])
    add_labeled(doc, "وظيفته", report["purpose"])
    add_labeled(doc, "مصدر البيانات", report["source"])
    add_labeled(doc, "شروط الإدراج", report["conditions"])
    add_labeled(doc, "طريقة الحساب", report["formula"])
    add_labeled(doc, "الفلاتر", report["filters"])
    add_labeled(doc, "ملاحظة", report["note"])
    divider = doc.add_paragraph()
    divider.paragraph_format.space_before = Pt(2)
    divider.paragraph_format.space_after = Pt(8)
    p_pr = divider._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E1EA")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

doc.add_page_break()
add_heading(doc, "5. ملاحظات رقابية وتشغيلية", 1)
add_status_box(doc, "هذا الدليل يشرح منطق النظام الحالي بعد المراجعة. أي تغيير في تعريف الربح أو سياسة العمولة يجب اعتماده إداريًا قبل تعديل المعادلات.", PALE_GOLD)
add_labeled(doc, "نطاق الربح", "الربح التشغيلي يطرح قيود المصروف فقط. إذا كان المطلوب صافي ربح شامل الرواتب والعمولات والمكافآت فيجب اعتماد ذلك كسياسة محاسبية منفصلة.")
add_labeled(doc, "التقارير المتشابهة", "الحضور الزائد والحضور المتكرر في اليوم يقيسان نفس الظاهرة الأساسية؛ الثاني يعرض كود العضو أيضًا.")
add_labeled(doc, "المكالمات", "تفاصيل إقناع المبيعات تربط المكالمات بالعملاء المحتملين عبر الهاتف، لذلك توحيد صيغة الهاتف ضروري للدقة.")
add_labeled(doc, "البيانات القديمة", "التدريب الشخصي يدعم شرط تصنيف private وشرط ربط الجلسات معًا حتى لا تختفي الاشتراكات القديمة.")
add_labeled(doc, "الباقات الطبية", "يلزم تصنيف الباقة medical في شاشة إعداد العضويات؛ الاسم وحده لم يعد كافيًا.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("report guide created")
